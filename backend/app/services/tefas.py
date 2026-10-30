"""
TEFAS fund prices from the "Fon Verileri" Excel export. Pure: no DB.

The export (one fund per file) is a single sheet:

    Rapor Bilgileri
    Dışa Aktarım Tarihi:   30.10.2026 15:19:13
    Toplam Kayıt Sayısı:   1258
    (blank)
    Tarih | Fiyat          ← header
    <date> | <price>       ← rows, any order; 0 before the fund existed

The fund code is not in the file (only in its name), so the caller
supplies it.
"""

import re
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

FUND_CODE = re.compile(r"^[A-Z0-9]{3}$")
HEADER = ("tarih", "fiyat")
# Header must appear within the first rows; guards against unrelated sheets.
MAX_HEADER_ROW = 20


class InvalidFundPrices(ValueError):
    """Rejected upload; the message is a code safe to show."""


def parse_price_export(content: bytes) -> dict[date, Decimal]:
    """Daily prices from a TEFAS Excel export, zero/blank rows dropped."""
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types
        raise InvalidFundPrices("not_excel") from exc

    try:
        rows = workbook.worksheets[0].iter_rows(values_only=True)
        for index, row in enumerate(rows):
            if index >= MAX_HEADER_ROW:
                raise InvalidFundPrices("no_price_table")
            if tuple(str(c).strip().lower() for c in row[:2]) == HEADER:
                break
        else:
            raise InvalidFundPrices("no_price_table")

        prices: dict[date, Decimal] = {}
        for row in rows:
            day, price = _parse_row(row)
            if day is not None and price is not None and price > 0:
                prices[day] = price
    finally:
        workbook.close()

    if not prices:
        raise InvalidFundPrices("no_prices")
    return prices


def _parse_row(row: tuple) -> tuple[date | None, Decimal | None]:
    if len(row) < 2 or row[0] is None or row[1] is None:
        return None, None
    raw_day, raw_price = row[0], row[1]

    if isinstance(raw_day, datetime):
        day = raw_day.date()
    elif isinstance(raw_day, date):
        day = raw_day
    else:
        try:
            day = datetime.strptime(str(raw_day).strip(), "%d.%m.%Y").date()
        except ValueError as exc:
            raise InvalidFundPrices("bad_row") from exc

    try:
        if isinstance(raw_price, (int, float)):
            price = Decimal(str(raw_price))
        else:  # Turkish text: 1.234,567890
            price = Decimal(str(raw_price).strip().replace(".", "").replace(",", "."))
    except InvalidOperation as exc:
        raise InvalidFundPrices("bad_row") from exc
    return day, price


def normalize_fund_code(code: str) -> str:
    normalized = code.strip().upper()
    if not FUND_CODE.match(normalized):
        raise InvalidFundPrices("invalid_fund_code")
    return normalized
